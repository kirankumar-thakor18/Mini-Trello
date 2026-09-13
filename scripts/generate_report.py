from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def center(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

# Title
t = doc.add_paragraph()
center(t)
run = t.add_run('Mini-Trello Kanban Board')
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0x00, 0x52, 0xCC)

sub = doc.add_paragraph()
center(sub)
r = sub.add_run('Project Report')
r.bold = True
r.font.size = Pt(16)

# Student details
doc.add_paragraph()
doc.add_heading('1. Student Details', level=1)
details = [
    ('Name', '[YOUR NAME]'),
    ('Enrollment Number', '[YOUR ENROLLMENT NUMBER]'),
    ('Semester', '7th Semester'),
    ('Subject', 'Mini Project'),
    ('Technology', 'Node.js + Express + SQLite + HTML/CSS/JS'),
]
for label, value in details:
    p = doc.add_paragraph()
    run = p.add_run(label + ': ')
    run.bold = True
    p.add_run(value)

# Overview
doc.add_heading('2. Project Overview', level=1)
doc.add_paragraph(
    'Mini-Trello is a simple single-page Kanban board application that allows users to create '
    'tasks and move them between three columns: To Do, In Progress, and Done. The project '
    'demonstrates practical application of Frontend development, Backend API creation, Database '
    'management, and Agile/Scrum methodology.'
)

doc.add_heading('3. Features', level=1)
for f in [
    'Create a new task (Title required, Description optional) - added to To Do',
    'View all tasks grouped by status in three columns',
    'Move tasks forward/backward using Next/Previous buttons',
    'Delete tasks with confirmation prompt',
    'Responsive, clean UI with distinct column colors',
    'All changes persist in the SQLite database',
]:
    doc.add_paragraph(f, style='List Bullet')

# Screenshots
doc.add_heading('4. Application Screenshots', level=1)

base = os.path.join(os.path.dirname(__file__), '..', 'screenshots')

def add_screenshot(path, caption):
    full = os.path.join(base, path)
    if os.path.exists(full):
        doc.add_picture(full, width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(caption)
        cr.italic = True
    else:
        doc.add_paragraph(caption + ' (screenshot missing)')

add_screenshot('board.png', 'Figure 1: Main 3-column Kanban board layout with task cards (Header with "+ Create New Task" button, To Do, In Progress, Done columns)')
add_screenshot('modal.png', 'Figure 2: "Create New Task" modal dialog')

# Agile Execution
doc.add_heading('5. Agile Execution Summary', level=1)
doc.add_paragraph(
    'The project was developed using the Agile/Scrum methodology and executed in two '
    'one-week sprints.'
)

doc.add_heading('Sprint 1 - Foundation & Setup', level=2)
for i in [
    'Conducted Sprint Planning and assigned tasks to team members.',
    'Designed the database schema for the tasks table (fields: id, title, description, status).',
    'Created backend API routes for Create and Read (GET all tasks, POST new task).',
    'Built the static HTML/CSS board layout with 3 columns.',
]:
    doc.add_paragraph(i, style='List Number')

doc.add_paragraph(
    'Deliverable at Sprint 1 Review: Working database, testable API endpoints (via Postman), '
    'and a static frontend UI.'
)

doc.add_heading('Sprint 2 - Integration & Delivery', level=2)
for i in [
    'Conducted Sprint 1 Retrospective and Sprint 2 Planning.',
    'Completed Update and Delete routes (PATCH status, DELETE task).',
    'Connected the frontend UI to the backend API using the fetch() method.',
    'Ensured card creation, movement, and deletion trigger correct API calls and update the UI dynamically.',
]:
    doc.add_paragraph(i, style='List Number')

doc.add_paragraph(
    'Deliverable at Sprint 2 Review: A fully functional Mini-Trello board where actions on the '
    'UI correctly persist in the database.'
)

doc.add_heading('User Stories Addressed', level=2)
stories = [
    ('HIGH', 'As a user, I want to create a new task so that I can add it to my board.'),
    ('HIGH', 'As a user, I want to see all my tasks grouped by status so I know what to work on.'),
    ('HIGH', 'As a user, I want to move a task from "To Do" to "In Progress" to "Done".'),
    ('MEDIUM', 'As a user, I want to delete a task if it was created by mistake.'),
]
for pr, story in stories:
    doc.add_paragraph(f'[{pr}] {story}', style='List Bullet')

# API Endpoints Table
doc.add_heading('6. API Endpoints Table', level=1)
doc.add_paragraph('Summary of developed REST endpoints:')

table = doc.add_table(rows=1, cols=4)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
headers = ['Method', 'Endpoint', 'Description', 'Status Code']
for i, h in enumerate(headers):
    p = hdr[i].paragraphs[0]
    r = p.add_run(h)
    r.bold = True

endpoints = [
    ('GET', '/api/tasks', 'Fetch all tasks', '200 OK'),
    ('POST', '/api/tasks', 'Create a new task (status=todo)', '201 Created'),
    ('PATCH', '/api/tasks/:id', 'Update task status/title/description', '200 OK'),
    ('PUT', '/api/tasks/:id', 'Replace task fields', '200 OK'),
    ('DELETE', '/api/tasks/:id', 'Delete a task', '204 No Content'),
]
for method, ep, desc, code in endpoints:
    row = table.add_row().cells
    row[0].text = method
    row[1].text = ep
    row[2].text = desc
    row[3].text = code

for r_ in table.rows:
    for c in r_.cells:
        for p in c.paragraphs:
            for run in p.runs:
                run.font.size = Pt(10)

doc.add_paragraph()
doc.add_heading('Sample Task Object (Reference)', level=2)
sample = doc.add_paragraph()
sample.add_run(
    '{\n'
    '  "id": "101",\n'
    '  "title": "Design Database Schema",\n'
    '  "description": "Create ER diagram for the task and user tables.",\n'
    '  "status": "in_progress"\n'
    '}'
)
sample.style = 'No Spacing'
for run in sample.runs:
    run.font.name = 'Consolas'
    run.font.size = Pt(10)

doc.add_heading('Sample Request / Response', level=3)
doc.add_paragraph('POST /api/tasks - Request body:')
req = doc.add_paragraph('{ "title": "Set up Express server", "description": "Create basic routes." }')
req.style = 'No Spacing'
for run in req.runs:
    run.font.name = 'Consolas'
    run.font.size = Pt(10)

doc.add_paragraph('Response (201 Created):')
res = doc.add_paragraph('{ "id": "7", "title": "Set up Express server", "description": "Create basic routes.", "status": "todo" }')
res.style = 'No Spacing'
for run in res.runs:
    run.font.name = 'Consolas'
    run.font.size = Pt(10)

# Team breakdown
doc.add_heading('7. Team Responsibility Breakdown', level=1)
team = doc.add_table(rows=1, cols=3)
team.style = 'Light Grid Accent 1'
hdr = team.rows[0].cells
for i, h in enumerate(['Team Member', 'Role', 'Responsibility']):
    p = hdr[i].paragraphs[0]
    r = p.add_run(h)
    r.bold = True
for member, role, resp in [
    ('[Member 1]', '[Role]', '[Database + Backend API]'),
    ('[Member 2]', '[Role]', '[Frontend Layout + Integration]'),
    ('[Member 3]', '[Role]', '[Testing + Documentation]'),
]:
    row = team.add_row().cells
    row[0].text = member
    row[1].text = role
    row[2].text = resp

doc.add_paragraph()
doc.add_paragraph(
    'Note: Replace all [bracketed] placeholders with actual student/team details before submission.'
)

out = os.path.join(base, '..', 'Mini-Trello_Project_Report.docx')
doc.save(out)
print('Saved:', os.path.abspath(out))
